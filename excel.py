from openpyxl import load_workbook
from itertools import groupby
import tracemalloc

tracemalloc.start()

def memoria():
    current, peak = tracemalloc.get_traced_memory()

    return (
        f"RAM tracemalloc: "
        f"{current / 1024**2:.1f} MB | "
        f"pico: {peak / 1024**2:.1f} MB"
    )

HEADERS = {
    'cod_emp': 'Empresa',
    'cod_amb': 'Ambito',
    'nro_liq': 'Liquidación',
    'des': 'Desc. Liquidación',
    'aaa_dev': 'Año',
    'mmm_dev': 'Mes',
    'fch_dde': 'Fecha Desde',
    'fch_hta': 'Fecha Hasta',
    'aaa_gan': 'Año Gan.',
    'mmm_gan': 'Mes Gan.',
    'nro_leg': 'Legajo',
    'ape_nom': 'Apellido Y Nombre',
    'fch_ini_r': 'Ingreso Relat.',
    'fch_ini': 'Ingreso',
    'fin_con': 'Fin Contrat.',
    'fch_egr': 'Egreso',
    'cod_tax': 'CUIT',
    'cod_cnv': 'Convenio',
    'lug_pag': 'Lugar De Pago',
    'adi_leg_1': 'Adi. Leg. 1',
    'den_adi_leg_1': 'Adic. Leg. 1',
    'adi_leg_2': 'Adi. Leg. 2',
    'den_adi_leg_2': 'Adic. Leg. 2',
    'adi_leg_3': 'Adi. Leg. 3',
    'den_adi_leg_3': 'Adic. Leg. 3',
    'adi_leg_4': 'Adi. Leg. 4',
    'den_adi_leg_4': 'Adic. Leg. 4',
    'adi_per_1': 'Adi. Per. 1',
    'den_adi_per_1': 'Adic. Per. 1',
    'adi_per_2': 'Adi. Per. 2',
    'den_adi_per_2': 'Adic. Per. 2',
    'cod_cpt': 'Concepto',
    'den_cpt': 'Den. Concepto',
    'den_rec': 'Den. Recibo',
    'den_aux': 'Den. Auxiliar',
    'cod_att': 'Agr. Imput.',
    'tip_cpt': 'Tipo Cpto.',
    'ord_cal': 'Ord. Calc.',
    'ord_rec': 'Ord.',
    'col_rec': 'Col. Recibo',
    'uni_med': 'U M',
    'can': 'Cantidad',
    'mto': 'Monto',
    'mto_rec': 'Mto. Recibo',
    'cod_are': 'Area',
    'den_are': 'Den. Area',
    'cod_cls': 'Clase',
    'den_cls': 'Den. Clase',
}


def load_data(path):
    print("load_data: inicio ", memoria())

    try:
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.worksheets[0]

        encabezados = [cell.value for cell in ws[5]]

        raw = []

        for row in ws.iter_rows(min_row=6, values_only=True):
            fila = dict(zip(encabezados, row))
            raw.append(fila)


        print("load_data: fin", memoria())

        return raw
    finally:
        wb.close()

def group_data(raw):
    print("group_data: inicio", memoria())

    gruped = []

    for legajo, grupo in groupby(raw, key=lambda row: row['nro_leg']):
        grupo = list(grupo)
        
        gruped.append({
            'legajo': legajo,
            'cuil': grupo[0]['cod_tax'],
            'dni': grupo[0]['cod_tax'],
            'nombre': grupo[0]['ape_nom'],
            'obra_social': [ 
                    row['den_rec'] for row in grupo if ('ret020' in str(row['cod_cpt']).lower())
                ],
            'ingreso': grupo[0]['fch_ini_r'],
            'categoria': grupo[0]['den_adi_leg_1'],
            'ingreso_reconocido': grupo[0]['fch_ini'],
            'egreso': grupo[0]['fch_egr'],
            'fin_contrato': grupo[0]['fin_con'],
            'tipo_liquidacion': grupo[0]['des'],
            'convenio': grupo[0]['cod_cnv'],
            'sueldo_bruto_categoria': [
                    row['mto'] for row in grupo if ('hcr008' in str(row['cod_cpt']).lower())
            ],
            'valor_bruto_categoria': [
                    row['mto'] for row in grupo if ('aux110' in str(row['cod_cpt']).lower())
            ],
            'periodo': grupo[0]['des'],
            'pago': f'{grupo[0]['lug_pag']} - 10/{grupo[0]['mmm_gan']}/{grupo[0]['aaa_gan']}',
            'fecha_pago_aportes': f'10/{grupo[0]['mmm_gan']}/{grupo[0]['aaa_gan']}',
            'conceptos': [
                {
                    'codigo': row['cod_cpt'],
                    'denominacion': row['den_cpt'],
                    'denominacion_recibo': row['den_rec'],
                    'denominacion_auxiliar': row['den_aux'],
                    'codigo_imputable': row['cod_att'],
                    'tipo_concepto': row['tip_cpt'],
                    'orden_calculo': row['ord_cal'],
                    'orden': row['ord_rec'],
                    'col_rec': row['col_rec'],
                    'um': row['uni_med'],
                    'cantidad': row['can'],
                    'monto': row['mto'],
                    'monto_recibo': row['mto_rec'],
                    'area': row['cod_are'],
                    'denominacion_area': row['den_are'],
                    'clase': row['cod_cls'],
                    'denominacion_clase': row['den_cls']
                }
                for row in grupo
            ]
        })

        print("group_data: nuevo grupo", memoria())
    print("group_data: fin", memoria())
    
    return gruped